"""Deck assembly reproduces the hand-approved slide pattern."""

from __future__ import annotations

import io
import sys
from pathlib import Path

import numpy as np
import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu, Inches

TOOL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(TOOL_ROOT))

from nifti_slides.deck import (  # noqa: E402
    BASELINE_TITLE,
    METRICS_SHAPE_NAME,
    POWERPOINT_BRIGHT20_LUT,
    DeckBuilder,
    SlideSpec,
    acceleration_label,
    case_label,
    comparison_title,
    format_caption,
    section_title,
    brightened_png_bytes,
)

TEMPLATE = TOOL_ROOT / "assets" / "slide_template.pptx"
A14 = "{http://schemas.microsoft.com/office/drawing/2010/main}"
DRAWING = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def test_section_title_matches_the_deck_wording() -> None:
    assert section_title("FMP_199", "MPRAGE_preGad") == "FMP 199 MPRAGE preGad"
    assert section_title("FMP_213", "MPRAGE_postGad") == "FMP 213 MPRAGE postGad"


def test_case_label_reads_the_variant_and_voxel_size() -> None:
    assert case_label("native_r3x2", (1.0, 1.0, 1.0)) == "1mm iso"
    assert case_label("native_r3x2", (0.8, 0.8, 0.8)) == "0.8mm iso"
    assert case_label("lr_x_1p5mm_r3x2", (1.5, 1.0, 1.0)) == "LR x 1.5mm"
    assert case_label("lr_xy_1p25mm_r3x2", (1.25, 1.25, 1.0)) == "LR xy 1.25mm"


def test_acceleration_label_keeps_the_lowercase_separator() -> None:
    assert acceleration_label("lr_x_1p5mm_r3x2") == "R3x2"
    assert acceleration_label("native_r4x1") == "R4x1"


def test_comparison_title_matches_the_hand_built_slides() -> None:
    assert comparison_title(
        "native_r3x2", "fista_r0", "Pre-SR", "", (1.0, 1.0, 1.0)
    ) == "Wave Fake R3x2 – 1mm iso, no reg, no SR"
    assert comparison_title(
        "lr_x_1p5mm_r3x2", "optimal_wavelet", "Post-SR", "shd", (1.5, 1.0, 1.0)
    ) == "Wave Fake R3x2 – LR x 1.5mm, wavelet, SHD"
    assert comparison_title(
        "lr_xy_1p25mm_r3x2", "fista_r0", "Post-SR", "cond_unet_retrain", (1.25, 1.25, 1.0)
    ) == "Wave Fake R3x2 – LR xy 1.25mm, no reg, retrained U-Net"


def test_format_caption_labels_the_mask() -> None:
    assert format_caption("head", 0.074336, 0.871965, 31.495612) == (
        " head mask     NRMSE 0.0743     SSIM 0.8720     PSNR 31.50 dB"
    )
    assert format_caption("brain", 0.1, 0.9, 30.0).startswith("brain mask")


def test_brightness_lut_is_monotone_and_brightens() -> None:
    values = np.asarray(POWERPOINT_BRIGHT20_LUT)
    assert values.shape == (256,)
    assert (np.diff(values) >= 0).all()
    assert values[0] == 0 and values[255] == 255
    # Midtones gain roughly the 1.25x PowerPoint applies for +20%.
    assert values[128] == pytest.approx(160, abs=2)


def _write_strip(path: Path, width: int = 64, height: int = 32) -> Path:
    ramp = np.linspace(0, 65535, width * height, dtype=np.uint16).reshape(height, width)
    Image.fromarray(ramp).save(path, format="TIFF", compression="tiff_lzw")
    return path


def test_brightened_png_bytes_bakes_brightness_without_touching_disk(
    tmp_path: Path,
) -> None:
    tiff = _write_strip(tmp_path / "strip.tiff")
    raster, width, height = brightened_png_bytes(tiff)

    assert (width, height) == (64, 32)
    assert isinstance(raster, bytes) and raster[:8] == b"\x89PNG\r\n\x1a\n"
    # The only files present are the strip we wrote; no PNG was persisted.
    assert [path.name for path in tmp_path.iterdir()] == ["strip.tiff"]

    with Image.open(io.BytesIO(raster)) as image:
        assert image.mode == "L"
        baked = np.array(image)

    source = np.rint(np.array(Image.open(tiff)).astype(np.float64) / 65535 * 255)
    expected = np.asarray(POWERPOINT_BRIGHT20_LUT, np.uint8)[source.astype(np.uint8)]
    np.testing.assert_array_equal(baked, expected)
    assert baked.mean() > source.mean()  # the raster really is brighter


def test_brightened_png_bytes_rejects_eight_bit_input(tmp_path: Path) -> None:
    path = tmp_path / "eight.tiff"
    Image.fromarray(np.zeros((4, 4), np.uint8)).save(path, format="TIFF")
    with pytest.raises(ValueError, match="16-bit"):
        brightened_png_bytes(path)


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="styling template is absent")
def test_deck_builder_reproduces_the_slide_pattern(tmp_path: Path) -> None:
    raster, _, _ = brightened_png_bytes(_write_strip(tmp_path / "strip.tiff", 640, 256))
    captions = (
        format_caption("head", 0.1, 0.9, 30.0),
        format_caption("brain", 0.2, 0.8, 28.0),
    )

    builder = DeckBuilder(TEMPLATE)
    builder.add_section("FMP 199 MPRAGE preGad")
    builder.add_picture_slide(
        SlideSpec(BASELINE_TITLE, raster, 640, 256, with_axes=False, captions=())
    )
    builder.add_picture_slide(
        SlideSpec(
            "Wave Fake R3x2 – LR x 1.5mm, no reg, SHD",
            raster,
            640,
            256,
            with_axes=True,
            captions=captions,
        )
    )
    out = tmp_path / "deck.pptx"
    builder.save(out)

    deck = Presentation(str(out))
    assert len(deck.slides) == 3  # the template's asset slide is dropped
    assert deck.slides[0].slide_layout.name == "Section Header"

    baseline, comparison = deck.slides[1], deck.slides[2]

    picture = [s for s in baseline.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE][0]
    assert picture.width == Inches(13.0)
    # Exactly centered on the slide.
    assert picture.left * 2 + picture.width == deck.slide_width
    assert picture.top * 2 + picture.height == deck.slide_height
    # Sent to back: children 0 and 1 are nvGrpSpPr and grpSpPr.
    assert picture._element is list(baseline.shapes._spTree)[2]

    blip = picture._element.blipFill.find(f"{DRAWING}blip")
    effect = blip.findall(f".//{A14}brightnessContrast")
    assert [node.get("bright") for node in effect] == ["20000"]

    assert not [s for s in baseline.shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP]
    assert not [s for s in baseline.shapes if s.name == METRICS_SHAPE_NAME]

    groups = [s for s in comparison.shapes if s.shape_type == MSO_SHAPE_TYPE.GROUP]
    assert len(groups) == 1
    captions = [s for s in comparison.shapes if s.name == METRICS_SHAPE_NAME]
    assert len(captions) == 1
    lines = captions[0].text_frame.text.splitlines()
    assert lines == [
        format_caption("head", 0.1, 0.9, 30.0),
        format_caption("brain", 0.2, 0.8, 28.0),
    ]
    # The caption clears the picture it sits under.
    assert captions[0].top >= picture.top + picture.height


@pytest.mark.skipif(not TEMPLATE.is_file(), reason="styling template is absent")
def test_deck_builder_gives_copied_groups_unique_shape_ids(tmp_path: Path) -> None:
    raster, _, _ = brightened_png_bytes(_write_strip(tmp_path / "strip.tiff"))

    builder = DeckBuilder(TEMPLATE)
    builder.add_picture_slide(
        SlideSpec("t", raster, 64, 32, with_axes=True, captions=("c",))
    )
    out = tmp_path / "deck.pptx"
    builder.save(out)

    slide = Presentation(str(out)).slides[0]
    ids = [
        int(node.get("id"))
        for node in slide.shapes._spTree.iter(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"
        )
    ]
    assert len(ids) == len(set(ids))
